import asyncio
import base64
import json
import re
from collections.abc import Awaitable, Callable
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from typing import Any
from urllib.parse import urlencode, urljoin
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pydantic import SecretStr
from selectolax.parser import HTMLParser

from trade_calendar.adapters.base import SourceAdapter
from trade_calendar.adapters.errors import (
    AuthenticationError,
    HttpStatusError,
    NetworkError,
    ParseError,
    StructureChangedError,
)
from trade_calendar.adapters.http import HttpFetcher
from trade_calendar.adapters.types import NormalizedEvent, RawPayload, SourceEvent
from trade_calendar.models.domain import DatePrecision, EventStatus, Importance
from trade_calendar.watchlist import WatchedCompany

SEOUL = ZoneInfo("Asia/Seoul")
TAIPEI = ZoneInfo("Asia/Taipei")
TOKYO = ZoneInfo("Asia/Tokyo")
FINNHUB_URL = "https://finnhub.io/api/v1/calendar/earnings"
LONGBRIDGE_EARNINGS_URL = (
    "https://open.longbridge.com/docs/market/calendar/earnings-calendar"
)
JPX_URL = "https://www.jpx.co.jp/listing/event-schedules/financial-announcement/"
KIND_URL = "https://kind.krx.co.kr/corpgeneral/irschedule.do"
TWSE_URL = "https://openapi.twse.com.tw/v1/opendata/t187ap04_L"
LongbridgeDocument = dict[str, Any]
LongbridgeRunner = Callable[[list[str], date, date], Awaitable[LongbridgeDocument]]


class LongbridgeEarningsAdapter(SourceAdapter):
    source_key = "longbridge_earnings"
    version = "1.0.0"
    allow_empty = True
    url = LONGBRIDGE_EARNINGS_URL

    def __init__(
        self,
        companies: list[WatchedCompany],
        today: Callable[[], date] | None = None,
        command_runner: LongbridgeRunner | None = None,
        cli_path: str = "longbridge",
    ) -> None:
        self.companies = [company for company in companies if company.longbridge_symbol]
        self.today = today or (lambda: datetime.now(UTC).date())
        self.command_runner = command_runner or self._run_calendar_command
        self.cli_path = cli_path

    async def fetch(self) -> RawPayload:
        start = self.today() - timedelta(days=31)
        end = self.today() + timedelta(days=180)
        symbols = list(dict.fromkeys(
            company.longbridge_symbol
            for company in self.companies
            if company.longbridge_symbol
        ))
        groups: list[dict[str, object]] = []
        queries: list[dict[str, object]] = []
        for offset in range(0, len(symbols), 10):
            batch = symbols[offset:offset + 10]
            document = await self.command_runner(batch, start, end)
            rows = document.get("list")
            if not isinstance(rows, list):
                raise StructureChangedError("Longbridge earnings calendar list is missing")
            valid_groups = [row for row in rows if isinstance(row, dict)]
            groups.extend(valid_groups)
            queries.append({"symbols": batch, "groups": len(valid_groups)})
        content = json.dumps({
            "provider": "Longbridge",
            "from": start.isoformat(),
            "to": end.isoformat(),
            "queries": queries,
            "list": groups,
        }, ensure_ascii=False).encode()
        return RawPayload(
            source_key=self.source_key,
            url=self.url,
            content=content,
            content_type="application/json",
        )

    async def _run_calendar_command(
        self,
        symbols: list[str],
        start: date,
        end: date,
    ) -> LongbridgeDocument:
        command = [self.cli_path, "finance-calendar", "report"]
        for symbol in symbols:
            command.extend(("--symbol", symbol))
        command.extend((
            "--start", start.isoformat(),
            "--end", end.isoformat(),
            "--count", "200",
            "--format", "json",
            "--lang", "en",
        ))
        try:
            process = await asyncio.create_subprocess_exec(
                *command,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
            )
        except OSError as exc:
            raise NetworkError("Longbridge CLI is unavailable") from exc
        try:
            stdout, stderr = await asyncio.wait_for(process.communicate(), timeout=60)
        except TimeoutError as exc:
            process.kill()
            await process.wait()
            raise NetworkError("Longbridge earnings command timed out") from exc
        if process.returncode != 0:
            error_text = stderr.decode(errors="replace").casefold()
            if "not authenticated" in error_text or "auth token" in error_text:
                raise AuthenticationError("Longbridge CLI authentication is unavailable")
            raise NetworkError(
                f"Longbridge earnings command failed with exit code {process.returncode}"
            )
        try:
            document = json.loads(stdout)
        except json.JSONDecodeError as exc:
            raise ParseError("invalid Longbridge earnings calendar JSON") from exc
        if not isinstance(document, dict):
            raise StructureChangedError("Longbridge earnings response is not an object")
        return document

    def parse(self, payload: RawPayload) -> list[SourceEvent]:
        try:
            document = json.loads(payload.content)
        except json.JSONDecodeError as exc:
            raise ParseError("invalid aggregated Longbridge earnings JSON") from exc
        groups = document.get("list") if isinstance(document, dict) else None
        if not isinstance(groups, list):
            raise StructureChangedError("aggregated Longbridge earnings list is missing")
        companies = {
            company.longbridge_symbol.upper(): company
            for company in self.companies
            if company.longbridge_symbol
        }
        events: dict[str, SourceEvent] = {}
        for group in groups:
            if not isinstance(group, dict) or not isinstance(group.get("infos"), list):
                continue
            for row in group["infos"]:
                if not isinstance(row, dict):
                    continue
                symbol = str(row.get("symbol") or "").strip().upper()
                company = companies.get(symbol)
                if company is None:
                    continue
                raw_ext = row.get("ext")
                ext: dict[str, Any] = raw_ext if isinstance(raw_ext, dict) else {}
                raw_financial = ext.get("financial_report")
                financial: dict[str, Any] = (
                    raw_financial if isinstance(raw_financial, dict) else {}
                )
                event_date = _as_date(ext.get("local_date")) or _date_from_text(
                    row.get("date")
                )
                if event_date is None:
                    continue
                try:
                    fiscal_year = int(financial.get("fiscal_year") or event_date.year)
                    quarter = int(financial.get("period") or 0)
                except (TypeError, ValueError):
                    fiscal_year = event_date.year
                    quarter = 0
                reference = _longbridge_reference_period(
                    financial,
                    str(row.get("content") or ""),
                    fiscal_year,
                    quarter,
                )
                source_id = f"longbridge:{company.key}:{reference.replace(' ', '-')}"
                market_time = str(
                    financial.get("market_time") or row.get("date_type") or ""
                ).strip().casefold()
                time_label = {
                    "pre": "before market open",
                    "before": "before market open",
                    "post": "after market close",
                    "after": "after market close",
                }.get(market_time)
                events[source_id] = SourceEvent(
                    source_event_id=source_id,
                    title=f"{company.name_en} Earnings Release",
                    local_date=event_date,
                    original_timezone=_market_timezone(company.market),
                    original_time_text=(
                        f"{event_date.isoformat()} {time_label}"
                        if time_label
                        else event_date.isoformat()
                    ),
                    reference_period=reference,
                    status_text="completed" if event_date < self.today() else "expected",
                    url=payload.url,
                    raw={
                        **row,
                        "company_key": company.key,
                        "canonical_ticker": company.ticker,
                        "query_symbol": symbol,
                    },
                )
        return list(events.values())

    def normalize(self, event: SourceEvent) -> NormalizedEvent:
        company = _company_for_event(event, self.companies)
        if event.local_date is None:
            raise StructureChangedError("Longbridge earnings event is missing a date")
        return _normalized_corporate_event(
            event,
            company,
            event_type="earnings_release",
            status=(
                EventStatus.COMPLETED
                if event.local_date < self.today()
                else EventStatus.EXPECTED
            ),
            source_url=event.url or self.url,
        )


class FinnhubEarningsAdapter(SourceAdapter):
    source_key = "finnhub_earnings"
    version = "1.0.0"

    def __init__(
        self,
        fetcher: HttpFetcher,
        api_key: SecretStr,
        companies: list[WatchedCompany],
        today: Callable[[], date] | None = None,
    ) -> None:
        self.fetcher = fetcher
        self.api_key = api_key
        self.companies = companies
        self.today = today or (lambda: datetime.now(UTC).date())

    async def fetch(self) -> RawPayload:
        start = self.today() - timedelta(days=31)
        end = self.today() + timedelta(days=120)
        headers = {
            "Accept": "application/json",
            "X-Finnhub-Token": self.api_key.get_secret_value(),
        }
        rows: list[dict[str, object]] = []
        queries: list[dict[str, object]] = []
        successful_requests = 0
        for company in self.companies:
            if not company.finnhub_symbol:
                continue
            query = urlencode({
                "from": start.isoformat(),
                "to": end.isoformat(),
                "symbol": company.finnhub_symbol,
            })
            try:
                payload = await self.fetcher.get(
                    self.source_key,
                    f"{FINNHUB_URL}?{query}",
                    headers,
                )
            except HttpStatusError as exc:
                if "403" not in str(exc):
                    raise
                queries.append({
                    "company_key": company.key,
                    "symbol": company.finnhub_symbol,
                    "error": "not entitled on current Finnhub plan",
                    "count": 0,
                })
                continue
            successful_requests += 1
            try:
                document = json.loads(payload.content)
            except json.JSONDecodeError as exc:
                raise ParseError(f"invalid Finnhub JSON for {company.finnhub_symbol}") from exc
            if not isinstance(document, dict):
                raise StructureChangedError("Finnhub earnings response is not an object")
            error = document.get("error")
            if error:
                queries.append({
                    "company_key": company.key,
                    "symbol": company.finnhub_symbol,
                    "error": str(error),
                    "count": 0,
                })
                continue
            calendar = document.get("earningsCalendar")
            if not isinstance(calendar, list):
                raise StructureChangedError(
                    f"Finnhub earningsCalendar missing for {company.finnhub_symbol}"
                )
            count = 0
            for value in calendar:
                if not isinstance(value, dict):
                    continue
                row = dict(value)
                row["_company_key"] = company.key
                row["_query_symbol"] = company.finnhub_symbol
                rows.append(row)
                count += 1
            queries.append({
                "company_key": company.key,
                "symbol": company.finnhub_symbol,
                "count": count,
            })
        if successful_requests == 0:
            raise HttpStatusError("Finnhub rejected every watchlist request")
        content = json.dumps({
            "from": start.isoformat(),
            "to": end.isoformat(),
            "queries": queries,
            "earningsCalendar": rows,
        }, ensure_ascii=False).encode()
        return RawPayload(
            source_key=self.source_key,
            url=f"{FINNHUB_URL}?from={start.isoformat()}&to={end.isoformat()}",
            content=content,
            content_type="application/json",
        )

    def parse(self, payload: RawPayload) -> list[SourceEvent]:
        try:
            document = json.loads(payload.content)
        except json.JSONDecodeError as exc:
            raise ParseError("invalid aggregated Finnhub earnings JSON") from exc
        rows = document.get("earningsCalendar") if isinstance(document, dict) else None
        if not isinstance(rows, list):
            raise StructureChangedError("aggregated Finnhub earningsCalendar is missing")
        companies = {company.key: company for company in self.companies}
        events: dict[str, SourceEvent] = {}
        for row in rows:
            if not isinstance(row, dict):
                continue
            company = companies.get(str(row.get("_company_key") or ""))
            if company is None:
                continue
            try:
                event_date = date.fromisoformat(str(row["date"]))
                fiscal_year = int(row.get("year") or event_date.year)
                quarter = int(row.get("quarter") or 0)
            except (KeyError, TypeError, ValueError):
                continue
            reference = _reference_period(fiscal_year, quarter)
            source_id = f"finnhub:{company.key}:{reference.replace(' ', '-')}"
            hour = str(row.get("hour") or "").strip().casefold()
            time_label = {
                "bmo": "before market open",
                "amc": "after market close",
                "dmh": "during market hours",
            }.get(hour)
            events[source_id] = SourceEvent(
                source_event_id=source_id,
                title=f"{company.name_en} Earnings Release",
                local_date=event_date,
                original_timezone=_market_timezone(company.market),
                original_time_text=(
                    f"{event_date.isoformat()} {time_label}"
                    if time_label
                    else event_date.isoformat()
                ),
                reference_period=reference,
                status_text="completed" if event_date < self.today() else "expected",
                url=payload.url,
                raw={
                    **{key: value for key, value in row.items() if not key.startswith("_")},
                    "company_key": company.key,
                    "canonical_ticker": company.ticker,
                    "query_symbol": row.get("_query_symbol"),
                },
            )
        return list(events.values())

    def normalize(self, event: SourceEvent) -> NormalizedEvent:
        company = _company_for_event(event, self.companies)
        if event.local_date is None:
            raise StructureChangedError("Finnhub earnings event is missing a date")
        return _normalized_corporate_event(
            event,
            company,
            event_type="earnings_release",
            status=(
                EventStatus.COMPLETED
                if event.local_date < self.today()
                else EventStatus.EXPECTED
            ),
            source_url=event.url or FINNHUB_URL,
        )


class JpxEarningsScheduleAdapter(SourceAdapter):
    source_key = "jpx_earnings_schedule"
    version = "1.0.0"
    allow_empty = True
    url = JPX_URL

    def __init__(self, fetcher: HttpFetcher, companies: list[WatchedCompany]) -> None:
        self.fetcher = fetcher
        self.companies = [company for company in companies if company.market == "JP"]

    async def fetch(self) -> RawPayload:
        listing = await self.fetcher.get(
            self.source_key,
            self.url,
            {"Accept": "text/html"},
        )
        tree = HTMLParser(listing.content)
        links = sorted({
            urljoin(listing.url, href)
            for node in tree.css("a")
            if (href := node.attributes.get("href"))
            and href.casefold().endswith(".xlsx")
        })
        if not links:
            raise StructureChangedError("JPX earnings schedule workbook links are missing")
        documents: list[dict[str, str]] = []
        for url in links[:12]:
            workbook = await self.fetcher.get(
                self.source_key,
                url,
                {"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
            )
            documents.append({
                "url": workbook.url,
                "content": base64.b64encode(workbook.content).decode("ascii"),
            })
        return RawPayload(
            source_key=self.source_key,
            url=listing.url,
            content=json.dumps({"documents": documents}).encode(),
            content_type="application/json",
        )

    def parse(self, payload: RawPayload) -> list[SourceEvent]:
        try:
            document = json.loads(payload.content)
        except json.JSONDecodeError as exc:
            raise ParseError("invalid aggregated JPX workbook payload") from exc
        workbooks = document.get("documents") if isinstance(document, dict) else None
        if not isinstance(workbooks, list):
            raise StructureChangedError("JPX workbook payload is missing documents")
        companies = {company.ticker: company for company in self.companies}
        events: dict[str, SourceEvent] = {}
        for item in workbooks:
            if not isinstance(item, dict):
                continue
            try:
                binary = base64.b64decode(str(item["content"]), validate=True)
                workbook = load_workbook(BytesIO(binary), read_only=True, data_only=True)
            except Exception as exc:
                raise ParseError("invalid JPX earnings schedule workbook") from exc
            sheet = workbook[workbook.sheetnames[0]]
            for row in sheet.iter_rows(min_row=6, values_only=True):
                if len(row) < 9:
                    continue
                ticker = _jpx_ticker(row[1])
                company = companies.get(ticker)
                scheduled = _as_date(row[0])
                fiscal_year_end = _as_date(row[4])
                period = _jpx_period(row[8])
                if (
                    company is None
                    or scheduled is None
                    or fiscal_year_end is None
                    or period is None
                ):
                    continue
                reference = (
                    f"FY{fiscal_year_end.year}"
                    if period == "FY"
                    else f"FY{fiscal_year_end.year} {period}"
                )
                source_id = f"jpx:{company.key}:{reference.replace(' ', '-')}"
                events[source_id] = SourceEvent(
                    source_event_id=source_id,
                    title=f"{company.name_en} Earnings Release",
                    local_date=scheduled,
                    original_timezone="Asia/Tokyo",
                    original_time_text=scheduled.isoformat(),
                    reference_period=reference,
                    status_text=(
                        "completed" if scheduled < datetime.now(TOKYO).date() else "confirmed"
                    ),
                    url=str(item.get("url") or payload.url),
                    raw={
                        "company_key": company.key,
                        "canonical_ticker": company.ticker,
                        "scheduled_date": scheduled.isoformat(),
                        "fiscal_year_end": fiscal_year_end.isoformat(),
                        "period": period,
                        "issue_name": str(row[3] or ""),
                    },
                )
            workbook.close()
        return list(events.values())

    def normalize(self, event: SourceEvent) -> NormalizedEvent:
        company = _company_for_event(event, self.companies)
        if event.local_date is None:
            raise StructureChangedError("JPX earnings event is missing a date")
        return _normalized_corporate_event(
            event,
            company,
            event_type="earnings_release",
            status=(
                EventStatus.COMPLETED
                if event.local_date < datetime.now(TOKYO).date()
                else EventStatus.CONFIRMED
            ),
            source_url=event.url or self.url,
        )


class KrxKindEarningsCallAdapter(SourceAdapter):
    source_key = "krx_kind_earnings_calls"
    version = "1.0.0"
    allow_empty = True
    url = KIND_URL

    def __init__(
        self,
        fetcher: HttpFetcher,
        companies: list[WatchedCompany],
        today: Callable[[], date] | None = None,
    ) -> None:
        self.fetcher = fetcher
        self.companies = [company for company in companies if company.market == "KR"]
        self.today = today or (lambda: datetime.now(SEOUL).date())

    async def fetch(self) -> RawPayload:
        start = self.today() - timedelta(days=7)
        end = self.today() + timedelta(days=120)
        query = urlencode({
            "method": "searchIRScheduleSub",
            "forward": "searchirschedule_sub",
            "currentPageSize": 3000,
            "pageIndex": 1,
            "marketType": 1,
            "fromDate": start.isoformat(),
            "toDate": end.isoformat(),
            "orderMode": 4,
            "orderStat": "D",
        })
        return await self.fetcher.get(
            self.source_key,
            f"{self.url}?{query}",
            {"Accept": "text/html"},
        )

    def parse(self, payload: RawPayload) -> list[SourceEvent]:
        tree = HTMLParser(payload.content)
        events: dict[str, SourceEvent] = {}
        for row in tree.css("tr"):
            cells = row.css("td")
            if len(cells) < 6:
                continue
            company_text = cells[1].text(separator=" ", strip=True)
            company = _company_by_name(company_text, self.companies)
            title = cells[2].text(separator=" ", strip=True)
            if company is None or not _is_earnings_related(title):
                continue
            try:
                event_date = date.fromisoformat(cells[4].text(strip=True))
            except ValueError:
                continue
            time_text = cells[5].text(strip=True)
            starts_at: datetime | None = None
            time_match = re.fullmatch(r"(\d{1,2}):(\d{2})", time_text)
            if time_match:
                local = datetime(
                    event_date.year,
                    event_date.month,
                    event_date.day,
                    int(time_match.group(1)),
                    int(time_match.group(2)),
                    tzinfo=SEOUL,
                )
                starts_at = local.astimezone(UTC)
            reference = _extract_reference_period(title, event_date.year)
            event_key = reference or event_date.isoformat()
            source_id = f"krx-kind:{company.key}:{event_key.replace(' ', '-')}"
            sequence = ""
            link = cells[2].css_first("a")
            if link is not None:
                sequence_match = re.search(
                    r"fnDetailView\(['\"]?(\d+)",
                    link.attributes.get("onclick") or "",
                )
                sequence = sequence_match.group(1) if sequence_match else ""
            events[source_id] = SourceEvent(
                source_event_id=source_id,
                title=f"{company.name_en} Earnings Call",
                starts_at=starts_at,
                local_date=None if starts_at else event_date,
                original_timezone="Asia/Seoul",
                original_time_text=(
                    f"{event_date.isoformat()} {time_text} Asia/Seoul"
                    if starts_at
                    else event_date.isoformat()
                ),
                reference_period=reference,
                status_text=(
                    "completed" if event_date < self.today() else "confirmed"
                ),
                url=payload.url,
                raw={
                    "company_key": company.key,
                    "canonical_ticker": company.ticker,
                    "kind_company_name": company_text,
                    "kind_title": title,
                    "place": cells[3].text(separator=" ", strip=True),
                    "event_date": event_date.isoformat(),
                    "event_time": time_text,
                    "ir_sequence": sequence,
                },
            )
        return list(events.values())

    def normalize(self, event: SourceEvent) -> NormalizedEvent:
        company = _company_for_event(event, self.companies)
        event_date = (
            event.starts_at.astimezone(SEOUL).date() if event.starts_at else event.local_date
        )
        if event_date is None:
            raise StructureChangedError("KIND earnings call is missing a date")
        return _normalized_corporate_event(
            event,
            company,
            event_type="earnings_call",
            status=(
                EventStatus.COMPLETED if event_date < self.today() else EventStatus.CONFIRMED
            ),
            source_url=event.url or self.url,
        )


class TwseEarningsCallAdapter(SourceAdapter):
    source_key = "twse_earnings_calls"
    version = "1.0.0"
    allow_empty = True
    url = TWSE_URL

    def __init__(self, fetcher: HttpFetcher, companies: list[WatchedCompany]) -> None:
        self.fetcher = fetcher
        self.companies = [company for company in companies if company.market == "TW"]

    async def fetch(self) -> RawPayload:
        return await self.fetcher.get(
            self.source_key,
            self.url,
            {"Accept": "application/json"},
        )

    def parse(self, payload: RawPayload) -> list[SourceEvent]:
        try:
            document = json.loads(payload.content)
        except json.JSONDecodeError as exc:
            raise ParseError("invalid TWSE material information JSON") from exc
        if not isinstance(document, list):
            raise StructureChangedError("TWSE material information response is not a list")
        companies = {company.ticker: company for company in self.companies}
        events: dict[str, SourceEvent] = {}
        for raw_row in document:
            if not isinstance(raw_row, dict):
                continue
            row = {str(key).strip(): value for key, value in raw_row.items()}
            company = companies.get(str(row.get("公司代號") or "").strip())
            subject = str(row.get("主旨") or "").strip()
            description = str(row.get("說明") or "").strip()
            combined = f"{subject}\n{description}"
            if company is None or "法人說明會" not in combined:
                continue
            reference = _extract_tw_reference_period(combined)
            if reference is None:
                continue
            event_date = _extract_tw_event_date(combined)
            if event_date is None:
                event_date = _parse_tw_date(str(row.get("事實發生日") or ""))
            if event_date is None:
                continue
            event_time = _extract_tw_event_time(combined)
            starts_at: datetime | None = None
            if event_time is not None:
                local = datetime(
                    event_date.year,
                    event_date.month,
                    event_date.day,
                    event_time[0],
                    event_time[1],
                    tzinfo=TAIPEI,
                )
                starts_at = local.astimezone(UTC)
            source_id = f"twse:{company.key}:{reference.replace(' ', '-')}"
            events[source_id] = SourceEvent(
                source_event_id=source_id,
                title=f"{company.name_en} Earnings Call",
                description=subject,
                starts_at=starts_at,
                local_date=None if starts_at else event_date,
                original_timezone="Asia/Taipei",
                original_time_text=(
                    f"{event_date.isoformat()} {event_time[0]:02d}:{event_time[1]:02d} Asia/Taipei"
                    if event_time is not None
                    else event_date.isoformat()
                ),
                reference_period=reference,
                status_text=(
                    "completed"
                    if event_date < datetime.now(TAIPEI).date()
                    else "confirmed"
                ),
                url=payload.url,
                raw={
                    **row,
                    "company_key": company.key,
                    "canonical_ticker": company.ticker,
                    "event_date": event_date.isoformat(),
                    "event_time": (
                        f"{event_time[0]:02d}:{event_time[1]:02d}"
                        if event_time is not None
                        else None
                    ),
                },
            )
        return list(events.values())

    def normalize(self, event: SourceEvent) -> NormalizedEvent:
        company = _company_for_event(event, self.companies)
        event_date = (
            event.starts_at.astimezone(TAIPEI).date() if event.starts_at else event.local_date
        )
        if event_date is None:
            raise StructureChangedError("TWSE earnings call is missing a date")
        return _normalized_corporate_event(
            event,
            company,
            event_type="earnings_call",
            status=(
                EventStatus.COMPLETED
                if event_date < datetime.now(TAIPEI).date()
                else EventStatus.CONFIRMED
            ),
            source_url=event.url or self.url,
        )


def _normalized_corporate_event(
    event: SourceEvent,
    company: WatchedCompany,
    *,
    event_type: str,
    status: EventStatus,
    source_url: str,
) -> NormalizedEvent:
    title_suffix = "财报电话会" if event_type == "earnings_call" else "财报发布"
    return NormalizedEvent(
        source_event_id=event.source_event_id,
        title_zh=f"{company.name_zh}{title_suffix}",
        title_original=event.title,
        institution=company.name_en,
        country_code=company.market,
        category="corporate",
        event_type=event_type,
        status=status,
        importance=Importance.HIGH,
        date_precision=(DatePrecision.MINUTE if event.starts_at else DatePrecision.DATE),
        starts_at=event.starts_at,
        local_date=event.local_date,
        original_timezone=event.original_timezone,
        original_time_text=event.original_time_text,
        reference_period=event.reference_period,
        market_tags=[company.market],
        tickers=[company.ticker],
        source_url=source_url,
        raw=event.raw,
    )


def _company_for_event(
    event: SourceEvent,
    companies: list[WatchedCompany],
) -> WatchedCompany:
    key = str(event.raw.get("company_key") or "")
    company = next((item for item in companies if item.key == key), None)
    if company is None:
        raise StructureChangedError(f"watchlist company missing for event: {key}")
    return company


def _company_by_name(text: str, companies: list[WatchedCompany]) -> WatchedCompany | None:
    candidate = _name_key(text)
    for company in companies:
        if any(
            _name_key(alias) in candidate or candidate in _name_key(alias)
            for alias in company.source_names
        ):
            return company
    return None


def _name_key(value: str) -> str:
    return re.sub(r"[\W_]+", "", value, flags=re.UNICODE).casefold()


def _reference_period(year: int, quarter: int) -> str:
    return f"FY{year} Q{quarter}" if 1 <= quarter <= 4 else f"FY{year}"


def _longbridge_reference_period(
    financial: dict[str, Any],
    content: str,
    fiscal_year: int,
    quarter: int,
) -> str:
    period_type = str(financial.get("period_type") or "").casefold()
    lowered = content.casefold()
    if "semi-annual" in lowered or period_type in {"saf", "half", "half_year"}:
        return f"FY{fiscal_year} H1"
    return _reference_period(fiscal_year, quarter)


def _market_timezone(market: str) -> str:
    return {
        "US": "America/New_York",
        "JP": "Asia/Tokyo",
        "KR": "Asia/Seoul",
        "CN": "Asia/Shanghai",
        "TW": "Asia/Taipei",
        "HK": "Asia/Hong_Kong",
    }.get(market, "UTC")


def _jpx_ticker(value: object) -> str:
    if isinstance(value, int):
        return f"{value:04d}"
    if isinstance(value, float) and value.is_integer():
        return f"{int(value):04d}"
    return str(value or "").strip().split(".", 1)[0].zfill(4)


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _date_from_text(value: object) -> date | None:
    match = re.search(r"20\d{2}[./-]\d{1,2}[./-]\d{1,2}", str(value or ""))
    if match is None:
        return None
    return _as_date(match.group().replace(".", "-"))


def _jpx_period(value: object) -> str | None:
    text = str(value or "").strip().casefold()
    if "first" in text or "1q" in text:
        return "Q1"
    if "second" in text or "2q" in text:
        return "Q2"
    if "third" in text or "3q" in text:
        return "Q3"
    if "full" in text or text in {"fy", "fiscal year"}:
        return "FY"
    return None


def _is_earnings_related(title: str) -> bool:
    lowered = title.casefold()
    return any(marker in lowered for marker in ("실적", "결산", "earnings", "results"))


def _extract_reference_period(text: str, fallback_year: int) -> str | None:
    patterns = (
        r"(?P<year>20\d{2})\s*년?\s*(?:제\s*)?(?P<quarter>[1-4])\s*분기",
        r"(?:FY\s*)?(?P<year>20\d{2})\s*Q(?P<quarter>[1-4])",
        r"Q(?P<quarter>[1-4])\s*(?:FY\s*)?(?P<year>20\d{2})",
        r"(?P<quarter>[1-4])Q\s*(?P<year>20\d{2})",
    )
    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return f"FY{int(match.group('year'))} Q{int(match.group('quarter'))}"
    quarter_match = re.search(r"(?:제\s*)?([1-4])\s*분기|Q([1-4])|([1-4])Q", text, re.I)
    if quarter_match:
        quarter = next(int(value) for value in quarter_match.groups() if value)
        return f"FY{fallback_year} Q{quarter}"
    return None


def _parse_tw_date(value: str) -> date | None:
    text = value.strip()
    match = re.search(
        r"(?P<year>\d{3,4})\s*(?:年|[/.-])\s*"
        r"(?P<month>\d{1,2})\s*(?:月|[/.-])\s*(?P<day>\d{1,2})",
        text,
    )
    if match is None:
        compact = re.fullmatch(r"(?P<year>\d{3,4})(?P<month>\d{2})(?P<day>\d{2})", text)
        match = compact
    if match is None:
        return None
    year = int(match.group("year"))
    if year < 1911:
        year += 1911
    try:
        return date(year, int(match.group("month")), int(match.group("day")))
    except ValueError:
        return None


def _extract_tw_event_date(text: str) -> date | None:
    match = re.search(r"召開法人說明會之日期\s*[：:]\s*([^\r\n]+)", text)
    return _parse_tw_date(match.group(1)) if match else None


def _extract_tw_event_time(text: str) -> tuple[int, int] | None:
    match = re.search(
        r"召開法人說明會之時間\s*[：:]\s*(\d{1,2})\s*時(?:\s*(\d{1,2})\s*分)?",
        text,
    )
    if match is None:
        return None
    hour = int(match.group(1))
    minute = int(match.group(2) or 0)
    return (hour, minute) if 0 <= hour <= 23 and 0 <= minute <= 59 else None


def _extract_tw_reference_period(text: str) -> str | None:
    year_match = re.search(r"(?P<year>20\d{2}|1\d{2})\s*年?", text)
    quarter_match = re.search(
        r"第?\s*(?P<quarter>[1-4一二三四])\s*(?:季|季度)|Q(?P<q>[1-4])",
        text,
        re.I,
    )
    if year_match is None or quarter_match is None:
        compact = re.search(r"(?P<year>20\d{2})\s*Q(?P<quarter>[1-4])", text, re.I)
        if compact is None:
            return None
        return f"FY{int(compact.group('year'))} Q{int(compact.group('quarter'))}"
    year = int(year_match.group("year"))
    if year < 1911:
        year += 1911
    raw_quarter = quarter_match.group("quarter") or quarter_match.group("q")
    fallback_quarter = int(raw_quarter) if raw_quarter.isdigit() else 0
    quarter = {"一": 1, "二": 2, "三": 3, "四": 4}.get(
        raw_quarter,
        fallback_quarter,
    )
    return f"FY{year} Q{quarter}" if quarter else None
